import csv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from exceptions import ExportError


def export_to_csv(posts: list[object], filepath: str) -> None:
    if not posts:
        raise ExportError("No post to export.")

    posts_data = [post.to_dict() for post in posts]
    fieldnames = list(posts_data[0].keys())

    try:
        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(posts_data)
    except (IOError, OSError) as e:
        raise ExportError(f"Failed to write csv: {e}")


def export_to_excel(posts: list[object], filepath: str) -> None:
    if not posts:
        raise ExportError("No post to export.")

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "HN Posts"

        posts_data = [post.to_dict() for post in posts]
        fieldnames = list(posts_data[0].keys())

        ws.append(fieldnames)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for post_dict in posts_data:
            ws.append(list(post_dict.values()))

        for col_cells in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col_cells)
            col_letter = get_column_letter(col_cells[0].column)
            ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

        wb.save(filepath)
    except (IOError, OSError) as e:
        raise ExportError(f"Failed to write excel: {e}")


def export_posts(
    posts: list[object], format_choice: str, csv_path: str, excel_path: str
) -> list[str]:
    if format_choice == "csv":
        export_to_csv(posts, csv_path)
        return [csv_path]
    elif format_choice == "excel":
        export_to_excel(posts, excel_path)
        return [excel_path]
    elif format_choice == "both":
        export_to_csv(posts, csv_path)
        export_to_excel(posts, excel_path)
        return [csv_path, excel_path]
    else:
        raise ValueError(f"Unrecognized export format: '{format_choice}'")
